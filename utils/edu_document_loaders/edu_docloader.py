# rag_qa/edu_document_loaders/edu_docloader.py
# 该脚本用于: 加载 Word 文档（含表格、图片 OCR 识别）

from typing import Iterator
from tqdm import tqdm
from docx.table import _Cell, Table
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx import ImagePart
from PIL import Image
from io import BytesIO
import numpy as np
import os

from langchain_core.documents import Document
from langchain_core.document_loaders import BaseLoader

# 导入项目内部的 OCR 模块（根据你的目录结构）
from utils.edu_document_loaders.edu_ocr import get_ocr

# 导入项目日志
from base.logger import logger


class OCRDOCLoader(BaseLoader):
    """
    从 Word 文档中提取文本，包含表格和图片 OCR 识别。
    支持 .docx 格式。
    """

    def __init__(self, filepath: str) -> None:
        """初始化加载器，传入文件路径"""
        self.filepath = filepath

    def lazy_load(self) -> Iterator[Document]:
        """懒加载，逐文档生成 Document 对象"""
        text = self.doc2text(self.filepath)
        if text:
            yield Document(page_content=text, metadata={"source": self.filepath})
        else:
            logger.warning(f"文档 {self.filepath} 提取内容为空")

    def doc2text(self, filepath: str) -> str:
        """
        从 Word 文档中提取文本，包括表格内容和图片 OCR 文字。
        """
        if not os.path.exists(filepath):
            logger.error(f"文件不存在: {filepath}")
            return ""

        try:
            # OCR 引擎惰性初始化：只有文档里真有图片时才加载（纯文本 docx 无需 rapidocr）
            ocr = None
            doc = DocxDocument(filepath)

            resp = ""  # 最终文本结果

            def iter_block_items(parent):
                """迭代文档中的所有段落和表格"""
                if isinstance(parent, DocxDocumentType):
                    parent_elm = parent.element.body
                elif isinstance(parent, _Cell):
                    parent_elm = parent._tc
                else:
                    raise ValueError("OCRDOCLoader 无法解析该父节点类型")

                for child in parent_elm.iterchildren():
                    if isinstance(child, CT_P):
                        yield Paragraph(child, parent)
                    elif isinstance(child, CT_Tbl):
                        yield Table(child, parent)

            # 创建进度条
            total_blocks = len(doc.paragraphs) + len(doc.tables)
            b_unit = tqdm(total=total_blocks, desc="OCRDOCLoader 处理进度", unit="块", disable=True)

            for i, block in enumerate(iter_block_items(doc)):
                b_unit.set_description(f"OCRDOCLoader 处理第 {i} 块")
                b_unit.refresh()

                # 处理段落
                if isinstance(block, Paragraph):
                    paragraph_text = block.text.strip()
                    if paragraph_text:
                        resp += paragraph_text + "\n"

                    # 处理段落中的图片
                    images = block._element.xpath('.//pic:pic')
                    for image in images:
                        for img_id in image.xpath('.//a:blip/@r:embed'):
                            part = doc.part.related_parts.get(img_id)
                            if isinstance(part, ImagePart):
                                try:
                                    img = Image.open(BytesIO(part._blob))
                                    if ocr is None:
                                        ocr = get_ocr()
                                    result, _ = ocr(np.array(img))
                                    if result:
                                        ocr_text = "\n".join([line[1] for line in result])
                                        resp += ocr_text + "\n"
                                except Exception as e:
                                    logger.debug(f"图片 OCR 失败: {e}")

                # 处理表格
                elif isinstance(block, Table):
                    for row in block.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                cell_text = paragraph.text.strip()
                                if cell_text:
                                    resp += cell_text + "\n"

                b_unit.update(1)

            b_unit.close()
            return resp

        except Exception as e:
            logger.error(f"解析 Word 文档失败: {e}")
            return ""


if __name__ == '__main__':
    import glob

    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(current_dir))
    ai_data_dir = os.path.join(project_root, 'rag_qa', 'data', 'ai_data')

    # 列出所有文件
    print("📂 ai_data 目录内容:")
    for file in os.listdir(ai_data_dir):
        print(f"  - {file}")

    # 测试 test.docx
    test_file = os.path.join(ai_data_dir, 'test.docx')
    if os.path.exists(test_file):
        loader = OCRDOCLoader(filepath=test_file)
        docs = loader.load()
        print(f"✅ test.docx 提取到 {len(docs[0].page_content)} 字符" if docs else "❌ 加载失败")

    # 动态查找中文 .docx 文件
    pattern = os.path.join(ai_data_dir, '人工智能*.docx')
    matching = glob.glob(pattern)
    if matching:
        chinese_file = matching[0]
        print(f"\n找到中文文档: {os.path.basename(chinese_file)}")
        loader = OCRDOCLoader(filepath=chinese_file)
        docs = loader.load()
        if docs:
            print(f"✅ 提取到 {len(docs[0].page_content)} 字符")
            print("=" * 50)
            print(docs[0].page_content[:300])  # 前300字符预览
    else:
        print("未找到中文 .docx 文件")